from io import BytesIO
from pathlib import Path
import re

from docx import Document
from pypdf import PdfReader

from app.domain.enums import DocumentKind
from app.domain.models import ExtractedDocument


MAX_UPLOAD_BYTES = 10 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class DocumentExtractionError(ValueError):
    """Raised when an uploaded document cannot be safely extracted."""


def extract_document(filename: str, payload: bytes) -> ExtractedDocument:
    safe_name = Path(filename).name
    extension = Path(safe_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentExtractionError(
            f"Unsupported file type: {extension or 'no extension'}. "
            "Supported types are .txt, .md, .pdf and .docx."
        )
    if not payload:
        raise DocumentExtractionError("Uploaded file is empty")
    if len(payload) > MAX_UPLOAD_BYTES:
        raise DocumentExtractionError("Uploaded file exceeds the 10 MB limit")

    if extension in {".txt", ".md"}:
        content = _decode_text(payload)
    elif extension == ".pdf":
        content = _extract_pdf(payload)
    else:
        content = _extract_docx(payload)

    content = content.strip()
    if not content:
        raise DocumentExtractionError("No readable text was found in the document")

    return ExtractedDocument(
        name=safe_name,
        suggested_kind=_suggest_kind(safe_name),
        content=content,
        size_bytes=len(payload),
    )


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentExtractionError("Text encoding is not supported")


def _extract_pdf(payload: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(payload))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise DocumentExtractionError("Unable to parse the PDF document") from exc


def _extract_docx(payload: bytes) -> str:
    try:
        document = Document(BytesIO(payload))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        tables = []
        for table in document.tables:
            for row in table.rows:
                tables.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join([*paragraphs, *tables])
    except Exception as exc:
        raise DocumentExtractionError("Unable to parse the DOCX document") from exc


def _suggest_kind(filename: str) -> DocumentKind | None:
    stem = Path(filename).stem
    tokens = {token for token in re.split(r"[^A-Z0-9]+", stem.upper()) if token}
    if "ICD" in tokens or any(token.startswith("INTERFACE") for token in tokens) or "接口" in stem:
        return DocumentKind.ICD
    if "AS" in tokens or any(token.startswith("ARCH") for token in tokens) or "架构" in stem:
        return DocumentKind.AS
    if "FS" in tokens or any(token.startswith("FUNCTION") for token in tokens) or "功能" in stem:
        return DocumentKind.FS
    return None
